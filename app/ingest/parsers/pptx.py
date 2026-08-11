"""PPTX 파서 (python-pptx).

슬라이드 번호를 locator에 남기고, 발표자 노트는 별도 조각으로 둔다.
회의 자료는 노트에 실제 논의 내용이 있는 경우가 많다.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from .base import (
    EMPTY,
    ENCRYPTED,
    MAX_CHARS,
    OK,
    UNSUPPORTED,
    DocMeta,
    ParseResult,
    Section,
    clip,
    failed,
)

PARSER = "pptx"


def parse(path: Path) -> ParseResult:
    if path.suffix.lower() == ".ppt":
        return failed(PARSER, "구형 PPT 형식 — 지원하지 않음", UNSUPPORTED)
    try:
        deck = Presentation(str(path))
    except Exception as exc:
        message = str(exc).lower()
        if "encrypted" in message or "password" in message:
            return failed(PARSER, "암호로 보호된 문서", ENCRYPTED)
        return failed(PARSER, f"프레젠테이션 열기 실패: {exc}")

    meta = _read_meta(deck)
    budget = [MAX_CHARS]
    out: list[Section] = []
    ordinal = 0

    for number, slide in enumerate(deck.slides, start=1):
        parts = [t for shape in slide.shapes for t in _shape_text(shape)]
        text = "\n".join(p for p in parts if p)
        if text.strip():
            body = clip(text, budget)
            if not body:
                break
            ordinal += 1
            out.append(Section("slide", ordinal, f"슬라이드 {number}", body))

        note = _notes(slide)
        if note:
            body = clip(note, budget)
            if not body:
                break
            ordinal += 1
            out.append(Section("note", ordinal, f"슬라이드 {number} 노트", body))

    if not out:
        return ParseResult(status=EMPTY, parser=PARSER, meta=meta)
    return ParseResult(status=OK, parser=PARSER, sections=out, meta=meta)


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if shape.shape_type is not None and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                out.append("\t".join(cells))
        return out
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            out.append(text)
    if shape.shape_type == 6:  # GROUP
        for child in shape.shapes:
            out.extend(_shape_text(child))
    return out


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""


def _read_meta(deck) -> DocMeta:
    props = deck.core_properties
    return DocMeta(
        title=(props.title or "").strip() or None,
        author=(props.author or "").strip() or None,
        created=props.created.isoformat() if props.created else None,
        modified=props.modified.isoformat() if props.modified else None,
    )
